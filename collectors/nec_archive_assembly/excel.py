"""Excel 파일을 셀 격자로 디코딩한다. **해석하지 않는다.**

nec_archive(대선)의 excel.py 와 별개다 — 이 수집기는 지역구가 **시트 이름**으로
나뉜 파일(17·18대)을 읽어야 해서 시트 이름 나열·이름으로 조회가 추가로 필요하다.
"""

from __future__ import annotations

from pathlib import Path

Grid = list[list[str]]


def read_grid(path: Path, *, sheet: int | str = 0) -> Grid:
    """한 시트를 문자열 격자로 읽는다. `sheet` 는 인덱스 또는 이름."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _read_xlsx(path, sheet)
    if suffix == ".xls":
        return _read_xls(path, sheet)
    raise ValueError(f"지원하지 않는 확장자다: {path.name} (.xls/.xlsx 만 읽는다)")


def sheet_names(path: Path) -> list[str]:
    """파일 안의 시트 이름 목록. 지역구가 시트 단위인 파일에서 대상을 고를 때 쓴다."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    if suffix == ".xls":
        import xlrd

        wb = xlrd.open_workbook(path, on_demand=True)
        return list(wb.sheet_names())
    raise ValueError(f"지원하지 않는 확장자다: {path.name}")


def _read_xlsx(path: Path, sheet: int | str) -> Grid:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if isinstance(sheet, str) else wb.worksheets[sheet]
        return [
            ["" if cell is None else str(cell) for cell in row]
            for row in ws.iter_rows(values_only=True)
        ]
    finally:
        wb.close()


def _read_xls(path: Path, sheet: int | str) -> Grid:
    import xlrd

    wb = xlrd.open_workbook(path)
    ws = wb.sheet_by_name(sheet) if isinstance(sheet, str) else wb.sheet_by_index(sheet)
    return [
        ["" if ws.cell_value(r, c) == "" else str(ws.cell_value(r, c)) for c in range(ws.ncols)]
        for r in range(ws.nrows)
    ]
