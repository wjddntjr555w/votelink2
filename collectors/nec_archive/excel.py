"""Excel 파일을 셀 격자로 디코딩한다. **해석하지 않는다.**

fetch 단계에서만 쓰인다. 컨테이너 포맷을 푸는 것은 파싱이 아니다
(C-002 의 CSV 디코딩과 같은 위치다). 격자를 raw 로 저장해 두면
`--reparse` 가 Excel 라이브러리 없이도 동작한다.

`.xls` 는 xlrd, `.xlsx` 는 openpyxl 이다. xlrd 2.x 는 `.xls` 전용이라 둘 다 필요하다.
"""

from __future__ import annotations

from pathlib import Path

Grid = list[list[str]]


def read_grid(path: Path, *, sheet: int = 0) -> Grid:
    """첫 시트를 문자열 격자로 읽는다. 값 변환·정리를 하지 않는다."""
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _read_xlsx(path, sheet)
    if suffix == ".xls":
        return _read_xls(path, sheet)
    raise ValueError(f"지원하지 않는 확장자다: {path.name} (.xls/.xlsx 만 읽는다)")


def _read_xlsx(path: Path, sheet: int) -> Grid:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[workbook.sheetnames[sheet]]
        return [
            ["" if cell is None else str(cell) for cell in row]
            for row in worksheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


def _read_xls(path: Path, sheet: int) -> Grid:
    import xlrd

    book = xlrd.open_workbook(path)
    worksheet = book.sheet_by_index(sheet)
    return [
        [str(worksheet.cell_value(r, c)) for c in range(worksheet.ncols)]
        for r in range(worksheet.nrows)
    ]
